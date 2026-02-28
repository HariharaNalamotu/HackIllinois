"""
Input nodes — the starting points of every pipeline.

These nodes receive uploaded files (or fetch from external APIs)
and produce typed NodeOutput dicts for downstream processing.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any, Callable

from .base import BaseNode
from ..types import NodeOutput

# Supported text-readable extensions
_TEXT_EXTS = {".txt", ".md", ".pdf", ".html", ".htm", ".docx", ".csv", ".json", ".jsonl"}
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".gif", ".webp"}
_AUDIO_EXTS = {".wav", ".mp3", ".flac", ".ogg", ".m4a", ".aac"}
_TABULAR_EXTS = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".jsonl"}


class TextInputNode(BaseNode):
    """
    Reads text from uploaded files.

    Supported: .txt, .md, .pdf, .html, .docx, .csv, .json, .jsonl
    Output: {"type": "text", "texts": list[str], "filenames": list[str]}
    """

    def execute(
        self,
        inputs: list[NodeOutput],
        files: list[str],
        ctx: dict[str, Any],
        log: Callable[[str], None],
    ) -> NodeOutput:
        import sys
        sys.path.insert(0, "/app")  # chunk.py is mounted here in Modal
        from chunk import read_file  # type: ignore[import]

        texts: list[str] = []
        filenames: list[str] = []

        for fp in files:
            p = Path(fp)
            if p.suffix.lower() not in _TEXT_EXTS:
                log(f"  [TextInput] skipping unsupported: {p.name}")
                continue
            try:
                text = read_file(p)
                if len(text.strip()) >= 10:
                    texts.append(text)
                    filenames.append(p.name)
                    log(f"  [TextInput] read {p.name} ({len(text)} chars)")
                else:
                    log(f"  [TextInput] skipped (too short): {p.name}")
            except Exception as e:
                log(f"  [TextInput] error reading {p.name}: {e}")

        if not texts:
            raise ValueError("TextInputNode: no usable text extracted from uploaded files.")

        return {"type": "text", "texts": texts, "filenames": filenames}


class ImageInputNode(BaseNode):
    """
    Accepts uploaded image files.

    Output: {"type": "image", "paths": list[str], "filenames": list[str]}
    """

    def execute(
        self,
        inputs: list[NodeOutput],
        files: list[str],
        ctx: dict[str, Any],
        log: Callable[[str], None],
    ) -> NodeOutput:
        paths: list[str] = []
        filenames: list[str] = []

        for fp in files:
            p = Path(fp)
            if p.suffix.lower() not in _IMAGE_EXTS:
                log(f"  [ImageInput] skipping non-image: {p.name}")
                continue
            paths.append(str(p))
            filenames.append(p.name)
            log(f"  [ImageInput] accepted {p.name}")

        if not paths:
            raise ValueError("ImageInputNode: no valid image files uploaded.")

        return {"type": "image", "paths": paths, "filenames": filenames}


class AudioInputNode(BaseNode):
    """
    Accepts uploaded audio files.

    Output: {"type": "audio", "paths": list[str], "filenames": list[str]}
    """

    def execute(
        self,
        inputs: list[NodeOutput],
        files: list[str],
        ctx: dict[str, Any],
        log: Callable[[str], None],
    ) -> NodeOutput:
        paths: list[str] = []
        filenames: list[str] = []

        for fp in files:
            p = Path(fp)
            if p.suffix.lower() not in _AUDIO_EXTS:
                log(f"  [AudioInput] skipping non-audio: {p.name}")
                continue
            paths.append(str(p))
            filenames.append(p.name)
            log(f"  [AudioInput] accepted {p.name}")

        if not paths:
            raise ValueError("AudioInputNode: no valid audio files uploaded.")

        return {"type": "audio", "paths": paths, "filenames": filenames}


class SpreadsheetInputNode(BaseNode):
    """
    Reads tabular data from .csv / .xlsx / .json / .jsonl files.

    Output: {"type": "tabular", "rows": list[dict], "columns": list[str]}
    """

    def execute(
        self,
        inputs: list[NodeOutput],
        files: list[str],
        ctx: dict[str, Any],
        log: Callable[[str], None],
    ) -> NodeOutput:
        all_rows: list[dict] = []
        columns: list[str] = []

        for fp in files:
            p = Path(fp)
            suffix = p.suffix.lower()

            try:
                if suffix in {".csv", ".tsv"}:
                    delimiter = "\t" if suffix == ".tsv" else ","
                    with open(p, newline="", encoding="utf-8-sig") as f:
                        reader = csv.DictReader(f, delimiter=delimiter)
                        rows = list(reader)
                        if rows:
                            columns = list(rows[0].keys())
                            all_rows.extend(rows)

                elif suffix in {".xlsx", ".xls"}:
                    import openpyxl  # type: ignore[import]
                    wb = openpyxl.load_workbook(p, data_only=True)
                    ws = wb.active
                    hdrs = [str(c.value) for c in next(ws.iter_rows(max_row=1))]
                    columns = hdrs
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        all_rows.append(dict(zip(hdrs, row)))

                elif suffix == ".json":
                    data = json.loads(p.read_text())
                    if isinstance(data, list) and data:
                        if isinstance(data[0], dict):
                            columns = list(data[0].keys())
                            all_rows.extend(data)
                        else:
                            columns = ["value"]
                            all_rows.extend({"value": v} for v in data)

                elif suffix == ".jsonl":
                    for line in p.read_text().splitlines():
                        line = line.strip()
                        if not line:
                            continue
                        obj = json.loads(line)
                        if not columns and isinstance(obj, dict):
                            columns = list(obj.keys())
                        all_rows.append(obj)

                log(f"  [SpreadsheetInput] read {p.name}: {len(all_rows)} rows")

            except Exception as e:
                log(f"  [SpreadsheetInput] error reading {p.name}: {e}")

        if not all_rows:
            raise ValueError("SpreadsheetInputNode: no rows extracted from uploaded files.")

        return {"type": "tabular", "rows": all_rows, "columns": columns}


class APIInputNode(BaseNode):
    """
    Fetches data from an external HTTP endpoint.

    Params: url, method (GET), headers ({}), body (""), data_path ("")
    Output: {"type": "text"/"tabular", ...} depending on response
    """

    def execute(
        self,
        inputs: list[NodeOutput],
        files: list[str],
        ctx: dict[str, Any],
        log: Callable[[str], None],
    ) -> NodeOutput:
        import httpx

        url        = self.p("url", "")
        method     = self.p("method", "GET").upper()
        headers    = self.p("headers", {})
        body       = self.p("body", "")
        data_path  = self.p("data_path", "")

        if not url:
            raise ValueError("APIInputNode: 'url' param is required.")

        log(f"  [APIInput] {method} {url}")

        resp = httpx.request(
            method=method,
            url=url,
            headers=headers,
            content=body.encode() if body else None,
            timeout=30,
        )
        resp.raise_for_status()

        try:
            data = resp.json()
        except Exception:
            return {"type": "text", "texts": [resp.text], "filenames": ["api_response.txt"]}

        # Navigate data_path (dot-separated keys)
        if data_path:
            for key in data_path.split("."):
                if isinstance(data, dict):
                    data = data.get(key, data)
                elif isinstance(data, list) and key.isdigit():
                    data = data[int(key)]

        # Return tabular or text depending on shape
        if isinstance(data, list) and data and isinstance(data[0], dict):
            cols = list(data[0].keys())
            return {"type": "tabular", "rows": data, "columns": cols}
        if isinstance(data, list):
            return {"type": "text", "texts": [str(x) for x in data], "filenames": []}
        if isinstance(data, str):
            return {"type": "text", "texts": [data], "filenames": []}

        return {"type": "text", "texts": [json.dumps(data)], "filenames": []}
